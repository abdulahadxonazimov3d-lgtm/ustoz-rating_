import os
from contextlib import asynccontextmanager
from io import BytesIO
from datetime import datetime
from urllib.parse import urlencode
from dotenv import load_dotenv
from fastapi import FastAPI, Request, Form, Depends, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, PlainTextResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import func
from aiogram import Bot, Dispatcher
from aiogram.types import Update
from aiogram.fsm.storage.memory import MemoryStorage
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from database import SessionLocal, init_db, AdminUser, Direction, Group, Student, Teacher, GroupTeacher, Rating, verify_password, hash_password
from bot_handlers import register_handlers
from seed import seed_default_data

load_dotenv()
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
WEBHOOK_BASE_URL = os.getenv("WEBHOOK_BASE_URL", "").rstrip("/")
WEBHOOK_SECRET = os.getenv("WEBHOOK_SECRET", "change-this-secret")
WEBHOOK_PATH = f"/telegram-webhook/{WEBHOOK_SECRET}"
AUTO_SEED = os.getenv("AUTO_SEED", "true").lower() == "true"
SESSION_SECRET = os.getenv("SESSION_SECRET", "local-session-secret-change-me")

bot = Bot(BOT_TOKEN) if BOT_TOKEN else None
dp = Dispatcher(storage=MemoryStorage())
register_handlers(dp)

@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    if AUTO_SEED:
        seed_default_data()
    if bot and WEBHOOK_BASE_URL:
        await bot.set_webhook(f"{WEBHOOK_BASE_URL}{WEBHOOK_PATH}", drop_pending_updates=True)
    yield
    if bot:
        await bot.session.close()

app = FastAPI(title="Ustoz Rating V2", lifespan=lifespan)
app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET, max_age=60*60*24*7, same_site="lax")
templates = Jinja2Templates(directory="templates")

def is_logged_in(request): return request.session.get("admin_logged_in") is True
def require_admin(request: Request):
    if not is_logged_in(request):
        raise HTTPException(status_code=303, headers={"Location": "/login"})
    return True

@app.exception_handler(HTTPException)
async def http_exception_handler(request, exc):
    if exc.status_code == 303 and "Location" in exc.headers:
        return RedirectResponse(exc.headers["Location"], status_code=303)
    return PlainTextResponse(str(exc.detail), status_code=exc.status_code)

def redir(url, msg="", error=""):
    q = {}
    if msg:
        q["msg"] = msg
    if error:
        q["error"] = error

    if not q:
        return RedirectResponse(url, status_code=303)

    separator = "&" if "?" in url else "?"
    return RedirectResponse(url + separator + urlencode(q), status_code=303)

@app.get("/")
def home(request: Request): return RedirectResponse("/admin" if is_logged_in(request) else "/login", status_code=303)

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse("login.html", {"request": request, "error": None})

@app.post("/login", response_class=HTMLResponse)
def login_post(request: Request, username: str = Form(...), password: str = Form(...)):
    db = SessionLocal()
    try:
        admin = db.query(AdminUser).first()
        if not admin or admin.username != username.strip() or not verify_password(password, admin.password_hash):
            return templates.TemplateResponse("login.html", {"request": request, "error": "Login yoki parol xato."}, status_code=401)
        request.session["admin_logged_in"] = True
        return RedirectResponse("/admin", status_code=303)
    finally:
        db.close()

@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)

@app.post(WEBHOOK_PATH)
async def telegram_webhook(request: Request):
    if not bot: raise HTTPException(status_code=500, detail="BOT_TOKEN sozlanmagan")
    update = Update.model_validate(await request.json(), context={"bot": bot})
    await dp.feed_update(bot, update)
    return {"ok": True}

@app.get("/set-webhook", response_class=PlainTextResponse)
async def set_webhook():
    if not bot: return "BOT_TOKEN sozlanmagan."
    if not WEBHOOK_BASE_URL: return "WEBHOOK_BASE_URL sozlanmagan. Lokal kompyuterda bu normal holat."
    url = f"{WEBHOOK_BASE_URL}{WEBHOOK_PATH}"
    await bot.set_webhook(url, drop_pending_updates=True)
    return f"Webhook o'rnatildi: {url}"

@app.get("/health", response_class=PlainTextResponse)
def health(): return "OK"

@app.get("/admin", response_class=HTMLResponse)
def admin(request: Request, _: bool = Depends(require_admin)):
    db=SessionLocal()
    try:
        stats = dict(directions=db.query(Direction).count(), groups=db.query(Group).count(), students=db.query(Student).count(), teachers=db.query(Teacher).count(), ratings=db.query(Rating).count(), bad=db.query(Rating).filter_by(has_bad_words=True).count())
        ranking = db.query(Teacher.full_name, Teacher.subject, func.avg(Rating.score).label("avg_score"), func.count(Rating.id).label("cnt")).join(Rating, Rating.teacher_id==Teacher.id).group_by(Teacher.id).order_by(func.avg(Rating.score).desc()).all()
        return templates.TemplateResponse("admin.html", {"request":request,"stats":stats,"ranking":ranking,"active":"dashboard"})
    finally: db.close()

@app.get("/admin/settings", response_class=HTMLResponse)
def settings_page(request: Request, msg: str="", error: str="", _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        return templates.TemplateResponse("settings.html", {"request":request,"admin":db.query(AdminUser).first(),"msg":msg,"error":error,"active":"settings"})
    finally: db.close()

@app.post("/admin/settings")
def update_settings(username: str=Form(...), current_password: str=Form(...), new_password: str=Form(""), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        admin=db.query(AdminUser).first()
        if not verify_password(current_password, admin.password_hash): return redir("/admin/settings", error="Joriy parol noto'g'ri.")
        admin.username=username.strip()
        if new_password.strip(): admin.password_hash=hash_password(new_password)
        db.commit()
        return redir("/admin/settings", msg="Sozlamalar saqlandi.")
    finally: db.close()

# Generic pages kept explicit
@app.get("/admin/directions", response_class=HTMLResponse)
def directions(request: Request, msg: str="", error: str="", _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        edit=request.query_params.get("edit")
        edit_item=db.query(Direction).filter_by(id=int(edit)).first() if edit and edit.isdigit() else None
        return templates.TemplateResponse("directions.html", {"request":request,"items":db.query(Direction).order_by(Direction.name).all(),"edit_item":edit_item,"msg":msg,"error":error,"active":"directions"})
    finally: db.close()

@app.post("/admin/directions")
def add_direction(name: str=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if not db.query(Direction).filter_by(name=name.strip()).first():
            db.add(Direction(name=name.strip())); db.commit()
        return redir("/admin/directions", msg="Yo'nalish saqlandi.")
    finally: db.close()

@app.post("/admin/directions/update")
def upd_direction(direction_id:int=Form(...), name:str=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        i=db.query(Direction).filter_by(id=direction_id).first()
        if i: i.name=name.strip(); db.commit()
        return redir("/admin/directions", msg="Yo'nalish yangilandi.")
    finally: db.close()

@app.post("/admin/directions/delete")
def del_direction(direction_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if db.query(Group).filter_by(direction_id=direction_id).first(): return redir("/admin/directions", error="Bu yo'nalishda guruh bor.")
        i=db.query(Direction).filter_by(id=direction_id).first()
        if i: db.delete(i); db.commit()
        return redir("/admin/directions", msg="Yo'nalish o'chirildi.")
    finally: db.close()

@app.get("/admin/groups", response_class=HTMLResponse)
def groups(request: Request, msg: str="", error: str="", _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        edit=request.query_params.get("edit")
        edit_item=db.query(Group).filter_by(id=int(edit)).first() if edit and edit.isdigit() else None
        return templates.TemplateResponse("groups.html", {"request":request,"items":db.query(Group).join(Direction).order_by(Direction.name,Group.name).all(),"directions":db.query(Direction).order_by(Direction.name).all(),"edit_item":edit_item,"msg":msg,"error":error,"active":"groups"})
    finally: db.close()

@app.post("/admin/groups")
def add_group(name:str=Form(...), direction_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if not db.query(Group).filter_by(name=name.strip(), direction_id=direction_id).first():
            db.add(Group(name=name.strip(), direction_id=direction_id)); db.commit()
        return redir("/admin/groups", msg="Guruh saqlandi.")
    finally: db.close()

@app.post("/admin/groups/update")
def upd_group(group_id:int=Form(...), name:str=Form(...), direction_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        i=db.query(Group).filter_by(id=group_id).first()
        if i: i.name=name.strip(); i.direction_id=direction_id; db.commit()
        return redir("/admin/groups", msg="Guruh yangilandi.")
    finally: db.close()

@app.post("/admin/groups/delete")
def del_group(group_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if db.query(Student).filter_by(group_id=group_id).first(): return redir("/admin/groups", error="Bu guruhda o'quvchilar bor.")
        db.query(GroupTeacher).filter_by(group_id=group_id).delete()
        i=db.query(Group).filter_by(id=group_id).first()
        if i: db.delete(i)
        db.commit()
        return redir("/admin/groups", msg="Guruh o'chirildi.")
    finally: db.close()

@app.get("/admin/students", response_class=HTMLResponse)
def students(
    request: Request,
    msg: str = "",
    error: str = "",
    last_group_id: int = 0,
    filter_group_id: int = 0,
    _: bool = Depends(require_admin)
):
    db = SessionLocal()
    try:
        edit = request.query_params.get("edit")
        edit_item = db.query(Student).filter_by(id=int(edit)).first() if edit and edit.isdigit() else None

        groups = db.query(Group).join(Direction).order_by(Direction.name, Group.name).all()

        query = db.query(Student).join(Group).join(Direction)

        if filter_group_id:
            query = query.filter(Student.group_id == filter_group_id)

        items = query.order_by(Direction.name, Group.name, Student.full_name).all()

        return templates.TemplateResponse("students.html", {
            "request": request,
            "items": items,
            "groups": groups,
            "edit_item": edit_item,
            "last_group_id": last_group_id,
            "filter_group_id": filter_group_id,
            "msg": msg,
            "error": error,
            "active": "students"
        })
    finally:
        db.close()

@app.post("/admin/students")
@app.post("/admin/students")
def add_student(full_name: str = Form(...), group_id: int = Form(...), _: bool = Depends(require_admin)):
    db = SessionLocal()
    try:
        full_name = full_name.strip()

        if not db.query(Student).filter_by(full_name=full_name, group_id=group_id).first():
            db.add(Student(full_name=full_name, group_id=group_id))
            db.commit()
            return redir(
                f"/admin/students?last_group_id={group_id}&filter_group_id={group_id}",
                msg="O'quvchi saqlandi."
            )

        return redir(
            f"/admin/students?last_group_id={group_id}&filter_group_id={group_id}",
            error="Bu o'quvchi shu guruhda mavjud."
        )
    finally:
        db.close()

@app.post("/admin/students/update")
def upd_student(student_id:int=Form(...), full_name:str=Form(...), group_id:int=Form(...), is_active:str=Form(None), reset_telegram:str=Form(None), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        i=db.query(Student).filter_by(id=student_id).first()
        if i:
            i.full_name=full_name.strip(); i.group_id=group_id; i.is_active=is_active=="on"
            if reset_telegram=="on": i.telegram_id=None; i.username=None
            db.commit()
        return redir("/admin/students", msg="O'quvchi yangilandi.")
    finally: db.close()

@app.post("/admin/students/delete")
def del_student(student_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        db.query(Rating).filter_by(student_id=student_id).delete()
        i=db.query(Student).filter_by(id=student_id).first()
        if i: db.delete(i)
        db.commit()
        return redir("/admin/students", msg="O'quvchi o'chirildi.")
    finally: db.close()

@app.get("/admin/teachers", response_class=HTMLResponse)
def teachers(request: Request, msg: str="", error: str="", _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        edit=request.query_params.get("edit")
        edit_item=db.query(Teacher).filter_by(id=int(edit)).first() if edit and edit.isdigit() else None
        return templates.TemplateResponse("teachers.html", {"request":request,"items":db.query(Teacher).order_by(Teacher.full_name).all(),"groups":db.query(Group).join(Direction).order_by(Direction.name,Group.name).all(),"links":db.query(GroupTeacher).join(Group).join(Direction).join(Teacher).order_by(Direction.name,Group.name,Teacher.full_name).all(),"edit_item":edit_item,"msg":msg,"error":error,"active":"teachers"})
    finally: db.close()

@app.post("/admin/teachers")
def add_teacher(full_name:str=Form(...), subject:str=Form(""), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if not db.query(Teacher).filter_by(full_name=full_name.strip()).first():
            db.add(Teacher(full_name=full_name.strip(), subject=subject.strip())); db.commit()
        return redir("/admin/teachers", msg="O'qituvchi saqlandi.")
    finally: db.close()

@app.post("/admin/teachers/update")
def upd_teacher(teacher_id:int=Form(...), full_name:str=Form(...), subject:str=Form(""), is_active:str=Form(None), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        i=db.query(Teacher).filter_by(id=teacher_id).first()
        if i: i.full_name=full_name.strip(); i.subject=subject.strip(); i.is_active=is_active=="on"; db.commit()
        return redir("/admin/teachers", msg="O'qituvchi yangilandi.")
    finally: db.close()

@app.post("/admin/teachers/delete")
def del_teacher(teacher_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if db.query(Rating).filter_by(teacher_id=teacher_id).first(): return redir("/admin/teachers", error="Bu o'qituvchida baholar bor.")
        db.query(GroupTeacher).filter_by(teacher_id=teacher_id).delete()
        i=db.query(Teacher).filter_by(id=teacher_id).first()
        if i: db.delete(i)
        db.commit()
        return redir("/admin/teachers", msg="O'qituvchi o'chirildi.")
    finally: db.close()

@app.post("/admin/link")
def link(group_id:int=Form(...), teacher_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        if not db.query(GroupTeacher).filter_by(group_id=group_id, teacher_id=teacher_id).first():
            db.add(GroupTeacher(group_id=group_id, teacher_id=teacher_id)); db.commit()
        return redir("/admin/teachers", msg="Biriktirish saqlandi.")
    finally: db.close()

@app.post("/admin/unlink")
def unlink(link_id:int=Form(...), _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        i=db.query(GroupTeacher).filter_by(id=link_id).first()
        if i: db.delete(i); db.commit()
        return redir("/admin/teachers", msg="Biriktirish olib tashlandi.")
    finally: db.close()

@app.get("/admin/ratings", response_class=HTMLResponse)
def ratings(request: Request, _: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        return templates.TemplateResponse("ratings.html", {"request":request,"items":db.query(Rating).order_by(Rating.created_at.desc()).all(),"active":"ratings"})
    finally: db.close()

@app.get("/admin/ratings/export")
def export_ratings(_: bool=Depends(require_admin)):
    db=SessionLocal()
    try:
        rows=db.query(Rating).order_by(Rating.created_at.desc()).all()
        wb=Workbook(); ws=wb.active; ws.title="Baholar"
        headers=["Sana","Yo'nalish","Guruh","O'quvchi","Telegram ID","Username","O'qituvchi","Fan/izoh","Ball","Fikr","Holat"]
        ws.append(headers)
        for r in rows:
            ws.append([r.created_at.strftime("%Y-%m-%d %H:%M:%S"), r.group.direction.name, r.group.name, r.student.full_name, r.student.telegram_id or "", ("@"+r.student.username) if r.student.username else "", r.teacher.full_name, r.teacher.subject or "", r.score, r.comment or "", "Haqoratli" if r.has_bad_words else "Toza"])
        for c in ws[1]:
            c.fill=PatternFill("solid", fgColor="1E3A8A"); c.font=Font(color="FFFFFF", bold=True)
        stream=BytesIO(); wb.save(stream); stream.seek(0)
        return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="baholar_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx"'})
    finally: db.close()

@app.get("/admin/seed", response_class=PlainTextResponse)
def seed_now(_: bool=Depends(require_admin)):
    seed_default_data()
    return "Ma'lumotlar tekshirildi."
